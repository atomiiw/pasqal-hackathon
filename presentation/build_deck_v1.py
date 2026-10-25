"""Build a seven-minute Track A deck. Performance fixtures are MOCK, never EMU results.

Run from repository root: .venv/Scripts/python.exe presentation/build_deck.py
Requires python-pptx, matplotlib, numpy and the repository's Pulser environment.
"""
from pathlib import Path
import sys, json, csv
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Circle
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from cvrp.instances import Instance, validate
from cvrp.qubo import build_qubo, variable_count_table
from cvrp.hardware import make_map, couplings
from cvrp.embedding.fidelity import fidelity, validate_layout
from cvrp.reference import cluster_cost, groups_to_assignment
from pulser import Register, Sequence, Pulse
from pulser.devices import AnalogDevice
from pulser.waveforms import InterpolatedWaveform, RampWaveform

OUT = ROOT/'presentation'
ASSETS = OUT/'charts'
ASSETS.mkdir(parents=True, exist_ok=True)
BG='#0B1426'; PANEL='#121F35'; INK='#F2F5FA'; MUTED='#A6B6CB'
GRID='#2A3A52'; MINT='#58D9B0'; BLUE='#75ACFF'; CORAL='#FF997B'; GOLD='#F2CA70'
COLORS=[MINT,BLUE,CORAL,GOLD]
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':12,'text.color':INK,
 'axes.labelcolor':MUTED,'xtick.color':MUTED,'ytick.color':MUTED,
 'axes.facecolor':PANEL,'figure.facecolor':PANEL,'axes.edgecolor':GRID,
 'grid.color':GRID,'axes.titleweight':'bold','axes.titlecolor':INK,
 'savefig.facecolor':PANEL,'legend.frameon':False})

def finish(fig,name):
    path=ASSETS/f'{name}.png'
    fig.savefig(path,dpi=190,bbox_inches='tight',pad_inches=.16)
    plt.close(fig)
    return path

def style(ax,grid=True):
    ax.spines[['top','right']].set_visible(False)
    if grid: ax.grid(axis='y',alpha=.7); ax.set_axisbelow(True)

instances={f'{t}{s}':Instance.load(ROOT/f'instances/{t}_seed{s}.json') for t in 'SML' for s in range(3)}
weights=json.loads((ROOT/'instances/weights.json').read_text())
for inst in instances.values():
    ok,msg=validate(inst,inst.certificate)
    assert ok,msg

def mapplot(ax,inst,groups=None,title=None,demands=False):
    xy=np.array(inst.coords)
    if groups is None:
        groups=[[i for i,k in enumerate(inst.certificate) if k==v] for v in range(inst.K)]
    for k,g in enumerate(groups):
        load=sum(inst.demands[i] for i in g)
        ax.scatter(*xy[g].T,c=COLORS[k],s=105,zorder=3,edgecolors=INK,linewidth=.5,
                   label=f'V{k+1}: {load}/{inst.Q}')
        for i in g:
            label=f'{chr(65+i)} ({inst.demands[i]})' if demands else chr(65+i)
            ax.annotate(label,xy[i],xytext=(5,5),textcoords='offset points',fontsize=10)
    ax.scatter(*inst.depot,marker='s',s=110,c=INK,zorder=5)
    ax.annotate('Depot',inst.depot,xytext=(7,-13),textcoords='offset points',fontsize=10)
    ax.set(xlim=(-5,108),ylim=(-5,108),xlabel='x / distance units',ylabel='y / distance units')
    ax.set_aspect('equal'); style(ax,False)
    if title: ax.set_title(title,loc='left',pad=14,fontsize=15)
    ax.legend(loc='lower left',fontsize=9,labelcolor=INK,handletextpad=.2,columnspacing=.7,ncol=2)

fig,ax=plt.subplots(figsize=(6.8,4.4),layout='constrained')
mapplot(ax,instances['M2'],instances['M2'].references['model_groups'],demands=True)
hero=finish(fig,'01_customer_clustering')

fig,axs=plt.subplots(1,3,figsize=(12,3.65),layout='constrained')
for ax,t in zip(axs,'SML'):
    inst=instances[t+'0']; mapplot(ax,inst,title=f'{t}  |  N={inst.N}, K={inst.K}, Q={inst.Q}')
    ax.text(.02,1.01,f'ρ={inst.rho:.2f}   •   dmax/Q={inst.dmax_ratio:.2f}',transform=ax.transAxes,color=MUTED,fontsize=10)
    ax.set_xlabel('x'); ax.set_ylabel('y')
tiers=finish(fig,'02_instance_tiers')

fig,ax=plt.subplots(figsize=(5.7,3.5),layout='constrained')
table=variable_count_table(); x=np.arange(3); w=.23
for off,field,label,color in [(-w,'onehot','Assignment-only',MUTED),(0,'onehot_plus_slack','+ Binary slack',BLUE),(w,'ours','Our encoding',MINT)]:
    vals=[r[field] for r in table]; bars=ax.bar(x+off,vals,w,label=label,color=color,zorder=3)
    ax.bar_label(bars,padding=4,color=INK,fontsize=12)
ax.axhline(80,ls='--',color=CORAL,lw=1.5); ax.text(-.4,83,'80-atom limit',color=CORAL,fontsize=10)
ax.set(xticks=x,xticklabels=['S','M','L'],ylabel='Binary variables = atoms',ylim=(0,98))
ax.legend(loc='upper left',fontsize=9,labelcolor=INK); style(ax)
atoms=finish(fig,'03_atom_savings')

# A legal schematic lattice, not an optimized embedding. Couplings are computed honestly.
inst=instances['M0']; wt=weights[inst.name]; q=build_qubo(inst,wt['lam'],wt['A'])
hmap=make_map(q.a,q.B,dmm=False,s_factor=.85)
coords=np.array([[6*(col+.5*(row%2)),6*np.sqrt(3)/2*row] for row in range(5) for col in range(6)],float)
coords-=coords.mean(axis=0)
validate_layout(coords,q.labels)
fm=fidelity(coords,hmap.target_U)
fig,(ax,sc)=plt.subplots(1,2,figsize=(10,4.05),layout='constrained',gridspec_kw={'width_ratios':[1,1.15]})
ax.add_patch(Circle((0,0),38,fill=False,ls='--',edgecolor=MUTED,lw=1))
for k in range(inst.K):
    ids=np.arange(k,q.n_vars,inst.K); ax.scatter(*coords[ids].T,c=COLORS[k],s=55,label=f'vehicle bit {k+1}')
ax.set(xlim=(-41,41),ylim=(-41,41),xlabel='x / µm',ylabel='y / µm'); ax.set_aspect('equal')
ax.set_title('Legal schematic register',loc='left',fontsize=14); style(ax,False)
ax.text(.02,.02,'30 atoms • spacing 6.0 µm',transform=ax.transAxes,fontsize=10,color=MINT)
iu=np.triu_indices(q.n_vars,1); target=hmap.target_U[iu]; actual=couplings(coords)[iu]; want=target>0
sc.scatter(target[want],actual[want],s=15,alpha=.65,c=BLUE,label='Target > 0')
sc.scatter(target[~want],actual[~want],s=13,alpha=.6,c=CORAL,label='Target = 0')
top=max(target.max(),actual.max())*1.08
sc.plot([0,top],[0,top],'--',c=MINT,lw=1.3,label='Ideal: y = x')
sc.set(xlim=(-1,top),ylim=(-1,top),xlabel='Target coupling / rad µs⁻¹',ylabel='Realized coupling / rad µs⁻¹')
sc.set_title('Mismatch remains visible',loc='left',fontsize=14); style(sc)
sc.legend(fontsize=9,labelcolor=INK,loc='upper right')
embedding=finish(fig,'04_embedding_diagnostic')

# Illustrative candidate pulse: validated by Pulser, not run through an emulator.
reg=Register({label:tuple(c) for label,c in zip(q.labels,coords)})
seq=Sequence(reg,AnalogDevice); seq.declare_channel('global','rydberg_global')
amp=InterpolatedWaveform(4000,[0,10,0]); det=RampWaveform(4000,-30,hmap.delta_global)
seq.add(Pulse(amp,det,0),'global'); AnalogDevice.validate_register(reg)
fig,axs=plt.subplots(2,1,figsize=(8,3.5),sharex=True,layout='constrained')
t=np.arange(4000)/1000
axs[0].plot(t,np.array(amp.samples),c=MINT,lw=2.5); axs[0].axhline(12.566370614,ls='--',c=CORAL,lw=1)
axs[0].set(ylabel='Ω / rad µs⁻¹',ylim=(-.6,14)); axs[0].text(2.8,12.9,'device max',fontsize=9,color=CORAL)
axs[1].plot(t,np.array(det.samples),c=BLUE,lw=2.5); axs[1].axhline(0,c=MUTED,lw=.6)
axs[1].set(ylabel='δ / rad µs⁻¹',xlabel='Time / µs',xlim=(0,4))
for ax in axs: style(ax)
pulse=finish(fig,'05_candidate_pulse')

# Explicit synthetic fixtures. These are design placeholders, not measured counts.
mock=[]
for ti,tier in enumerate('SML'):
    for seed in range(3):
        for method in ['Baseline pulse','Tuned pulse']:
            tuned=method=='Tuned pulse'
            feasible=([.94,.73,.49] if tuned else [.64,.38,.16])[ti]+[-.03,.01,.02][seed]
            mg=([[0,0,0],[0,1.6,15],[11,15,19]] if tuned else [[0,4,9],[10,17,28],[24,32,43]])[ti][seed]
            cg=([[2,4,6],[6,10,18],[15,20,25]] if tuned else [[8,10,14],[16,23,32],[30,40,50]])[ti][seed]
            hit=([[.25,.18,.21],[.06,0,0],[0,0,0]] if tuned else [[.08,.04,0],[0,0,0],[0,0,0]])[ti][seed]
            mock.append(dict(data_status='MOCK_NOT_EMULATOR_OUTPUT',instance=f'{tier}_seed{seed}',method=method,seed=seed,
                shots=1000,feasibility_rate=round(feasible,3),model_gap_pct=mg,cvrp_gap_pct=cg,p_reach_model_ref=hit,
                tolerance_relative=1e-6,atoms=[4,30,60][ti]))
with (OUT/'mock_metrics.csv').open('w',newline='',encoding='utf-8') as f:
    writer=csv.DictWriter(f,fieldnames=list(mock[0])); writer.writeheader(); writer.writerows(mock)
fig,axs=plt.subplots(1,3,figsize=(12,3.1),layout='constrained')
for ax,key,title,ylim in zip(axs,['feasibility_rate','model_gap_pct','cvrp_gap_pct'],['Feasibility rate ↑','Model gap ↓','CVRP gap ↓'],[(0,105),(0,48),(0,55)]):
    for method,col,shift in [('Baseline pulse',BLUE,-.05),('Tuned pulse',MINT,.05)]:
        means=[]; lows=[]; highs=[]
        for i,tier in enumerate('SML'):
            vals=np.array([r[key] for r in mock if r['instance'].startswith(tier) and r['method']==method])
            if key=='feasibility_rate': vals*=100
            means.append(vals.mean()); lows.append(vals.mean()-vals.min()); highs.append(vals.max()-vals.mean())
            ax.scatter(np.full(3,i+shift),vals,color=col,s=15,alpha=.55)
        ax.errorbar(np.arange(3)+shift,means,yerr=[lows,highs],color=col,lw=2,marker='o',capsize=4,label=method)
    ax.set(xticks=np.arange(3),xticklabels=list('SML'),ylim=ylim,ylabel='%')
    ax.set_title(title,loc='left',fontsize=15); style(ax)
    if key!='feasibility_rate': ax.axhline(0,color=GOLD,ls='--',lw=1)
axs[0].legend(fontsize=9,labelcolor=INK,loc='lower left')
results=finish(fig,'06_mock_results')
fig,ax=plt.subplots(figsize=(5.5,1.5),layout='constrained')
shots=[1,10,35,80,160,300,550,1000]
ax.step(shots,[35,35,23,23,13,13,8,8],where='post',color=BLUE,label='Baseline pulse',lw=1.8)
ax.step(shots,[32,24,24,12,6,3,0,0],where='post',color=MINT,label='Tuned pulse',lw=1.8)
ax.set(xlabel='Cumulative shots',ylabel='Best gap / %',xlim=(0,1000),ylim=(-2,40)); style(ax)
trace=finish(fig,'06_mock_shot_trace')

case=instances['M2']; casew=weights[case.name]; refgroups=case.references['model_groups']; qgroups=casew['groups']
ca=groups_to_assignment(refgroups,case.N); qa=groups_to_assignment(qgroups,case.N)
refcost=cluster_cost(ca,case.dist_matrix()); qcost=cluster_cost(qa,case.dist_matrix())
assert abs(qcost-casew['model_cost'])<1e-6
assert validate(case,ca)[0] and validate(case,qa)[0]
loss=100*(qcost/refcost-1)
pen=lambda groups: sum((sum(case.demands[i] for i in g)-case.D/case.K)**2 for g in groups)
fig,axs=plt.subplots(1,2,figsize=(8.5,3.45),layout='constrained')
mapplot(axs[0],case,refgroups,title=f'Constrained optimum: C = {refcost:.2f}')
mapplot(axs[1],case,qgroups,title=f'Saved QUBO solution: C = {qcost:.2f}')
for ax,g in zip(axs,[refgroups,qgroups]):
    ax.text(.02,1.01,f'Balance penalty = {pen(g):.2f}',transform=ax.transAxes,color=MUTED,fontsize=10)
failure=finish(fig,'07_actual_balance_bias')

prs=Presentation(); prs.slide_width=Inches(16); prs.slide_height=Inches(9)
def rgb(s): return RGBColor.from_string(s.lstrip('#'))
def box(slide,x,y,w,h,fill=PANEL,line=None):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=rgb(fill)
    shape.line.fill.background() if line is None else None
    return shape
def text(slide,x,y,w,h,value,size=20,color=INK,bold=False,font='Aptos'):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame
    tf.word_wrap=True; tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.name=font
        p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color); p.space_after=Pt(8)
    return sh
def picture(slide,path,x,y,w,h):
    from PIL import Image
    iw,ih=Image.open(path).size; factor=min(w/iw,h/ih); nw,nh=iw*factor,ih*factor
    slide.shapes.add_picture(str(path),Inches(x+(w-nw)/2),Inches(y+(h-nh)/2),width=Inches(nw),height=Inches(nh))
def base(num,kicker,title,sub,status='REAL MODEL / INSTANCE DATA'):
    slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=rgb(BG)
    text(slide,.65,.35,10,.3,f'TRACK A  /  {kicker.upper()}',12,MINT,True)
    text(slide,.65,.92,14.7,.65,title,32,INK,True)
    text(slide,.65,1.7,14.7,.55,sub,17,MUTED)
    text(slide,.65,8.5,11,.25,f'DuQuantum × Pasqal  •  {status}',10,CORAL if 'MOCK' in status else MUTED)
    text(slide,14.25,8.5,1,.25,f'{num:02d} / 07',11,MUTED)
    return slide
def note(slide,value): slide.notes_slide.notes_text_frame.text=value
def caption(slide,value): text(slide,.75,7.78,14.5,.5,value,16,MUTED)

s=base(1,'The idea','Fewer atoms. Capacity-aware clusters.','A slack-free QUBO for customer-to-vehicle assignment on neutral-atom emulators.')
box(s,.65,2.55,5.2,4.85); text(s,.95,2.9,4.65,1.2,'One variable\n= one atom',34,MINT,True)
text(s,.95,4.38,4.55,1.2,'4 / 30 / 60 atoms\nfor S / M / L',27,INK,True)
text(s,.95,6.05,4.6,1.0,'Trade-off: load balancing is a surrogate for capacity.',19,CORAL)
box(s,6.1,2.55,9.25,4.85); picture(s,hero,6.35,2.7,8.7,4.5)
caption(s,'Shown: M_seed2 constrained clustering reference • labels = customer (demand); legend = vehicle load / capacity.')
note(s,'[35 秒] 我们做的是 Track A，capacity-constrained clustering。每个 customer 必须恰好分配给一辆 vehicle，每辆车的 load 不超过 capacity。我们用 slack-free QUBO，将 S 的变量数降到 4，M 和 L 分别是 30 和 60。右图是实际 instance 的 constrained reference，不是 quantum output。核心 trade-off 是：省掉 slack bits，但用 load balancing 代替严格的 capacity inequality。整个 deck 是展示草稿，后面的 performance 明确使用 mock data。')

s=base(2,'Instances','Three scales. Certified feasible.','Fixed depot (50, 50) • Euclidean distances • seeds 0, 1, 2 for each tier')
box(s,.65,2.55,14.7,4.6); picture(s,tiers,.85,2.65,14.3,4.35)
text(s,.95,7.26,14,.45,'9 instances  →  assignment certificate  →  capacity validator  →  exact classical references',20,MINT,True)
caption(s,'Maps show seed 0 certificates, not optimized outputs. Generator accepts 0.70 ≤ load ratio ρ ≤ 0.95.')
note(s,'[45 秒] S、M、L 是比赛固定的 instance tiers。S 是 4 customers、2 vehicles、capacity 7；M 是 10、3、20；L 是 15、4、30。这里展示每个 tier 的 seed 0 feasibility certificate。我们固定 depot 为 50,50，使用 Euclidean distances，三个 seeds 共 9 个 instances。生成器只接受 load ratio 在 0.70 到 0.95 之间的 instances，并检查 capacity 确实影响 unconstrained clustering。每个 instance 附一个合法 assignment，reference 用 classical exact methods 计算。')

s=base(3,'Formulation','Remove slack. Keep the cost quadratic.','Implemented in cvrp/qubo.py • positive pair couplings • fewer auxiliary variables')
box(s,.65,2.55,7.05,4.95); box(s,7.95,2.55,7.4,4.95)
text(s,.95,2.84,6.45,.45,'E = C + A P_assign + λ P_balance',23,MINT,True)
text(s,.95,3.55,6.45,.65,'C = Σₖ Σᵢ<ⱼ distanceᵢⱼ xᵢₖ xⱼₖ',20)
text(s,.95,4.23,6.45,.6,'P_assign = Σᵢ (1 − Σₖ xᵢₖ)²',20)
text(s,.95,4.88,6.45,.6,'P_balance = Σₖ (Lₖ − D/K)²',20)
text(s,.95,5.63,6.45,.7,'S: one bit/customer; assignment automatic.\nM/L: one-hot; assignment penalty required.',17,MUTED)
text(s,.95,6.6,6.45,.5,'Balance encourages capacity; it does not enforce it.',17,CORAL,True)
picture(s,atoms,8.1,2.85,7.1,4.35)
caption(s,'A and λ are instance-tuned in the current code; exact classical calibration is preprocessing, not quantum work.')
note(s,'[100 秒] 我们的 clustering objective 是同一 vehicle 内 customer pairwise distances 的总和。第一项希望同组 customers 接近；assignment penalty 要求每个 customer 恰好一个 vehicle；balance penalty 希望 load 靠近平均值 D/K。S 因为只有两辆 vehicle，每个 customer 一个 bit，0 和 1 就代表两辆车，因此 assignment 自动满足；S 的 objective 使用 same-cluster indicator，具体展开在 qubo.py。M 和 L 使用 one-hot，需要 assignment penalty。我们的 capacity term 没有 Q 的阈值，它是 surrogate，不是精确 inequality。右边 atom count 是实际代码算出来的：standard slack 是 14,45,80，我们是 4,30,60。A 和 lambda 是 instance-specific，当前 calibration 使用 exact classical solvers，这部分计算成本必须如实说明。')

s=base(4,'Embedding','Legal geometry ≠ faithful QUBO.','Vᵢⱼ = C₆ / rᵢⱼ⁶ • positive interactions • all pairs interact • shared global detuning',status='ILLUSTRATIVE REGISTER / COMPUTED DIAGNOSTICS')
box(s,.65,2.55,14.7,4.6); picture(s,embedding,.85,2.66,14.3,4.35)
text(s,.95,7.23,14,.4,f'Pulser-valid schematic • min spacing {fm["min_dist"]:.1f} µm • max radius {fm["max_radius"]:.1f} µm • coupling error η={fm["eta"]:.2f}',18,MINT)
caption(s,'η = ‖U − sB‖₂ / ‖sB‖₂ over i<j, including target-zero pairs. Layout is illustrative, not an optimized embedding.')
note(s,'[60 秒] 这是一个真实通过 Pulser register validation 的 30-atom schematic lattice，不是我们 embedder 的优化结果。所有 plotted couplings 都由这些 positions 用 C6/r6 计算。右图横轴是 M_seed0 的 scaled target，纵轴是 realized interactions；理想是 y=x。Target 为零的 pair 在硬件里仍然有 positive coupling，图中橙色点明确展示这种误差。我们用全部 pairs 的 relative L2 error，不会丢掉 zero-target pairs。Global detuning 只能提供一个 linear coefficient。M/L 的 linear mismatch 在 valid one-hot assignments 上是常数，但仍影响 invalid shots；S 不适用这个简化。这里说明 diagnostics 该怎么做，不声称已经得到优秀 embedding。')

s=base(5,'Pulse design','A device-valid starting pulse.','Illustrative global-only candidate • M_seed0 scaling • not executed on an emulator',status='ILLUSTRATIVE PULSE / PULSER-VALIDATED')
box(s,.65,2.55,10.1,4.95); picture(s,pulse,.9,2.8,9.6,4.45)
box(s,11,2.55,4.35,4.95)
text(s,11.3,2.9,3.7,.5,'4,000 ns',30,MINT,True)
text(s,11.3,3.53,3.7,.5,'of 6,000 ns allowed',17,MUTED)
text(s,11.3,4.28,3.7,.5,'Ω peak = 10 rad/µs',18)
text(s,11.3,4.95,3.7,.7,f'δ: −30 → {hmap.delta_global:.1f}\nrad/µs',19)
text(s,11.3,6.03,3.7,1.0,'Execution settings pending:\nemulator, shots, seeds,\nMPS bond dimension.',16,CORAL)
caption(s,'Final δ comes from make_map(..., dmm=False, s_factor=0.85). Valid pulse construction does not establish performance.')
note(s,'[45 秒] 我们把 pulse design 拆成 Rabi frequency 和 detuning 两条曲线。这个 candidate 总时长 4000 ns，Omega 从零升到 10，再回到零，detuning 从负 30 扫到图中的 final detuning。Final detuning 来自当前 global hardware map。我们实际用 Pulser 构建了 sequence，所有数值在 device limits 内。但这不是 emulator execution，不能把合法的 pulse 等同于成功的优化。正式结果必须补 emulator、shots、seeds，如果使用 emu-mps，还要给 bond dimension。')

s=base(6,'Performance storyboard','Measure feasibility AND solution quality.','MOCK DATA — synthetic plotting fixtures, not quantum runs, forecasts, or measured improvements',status='MOCK PERFORMANCE — REPLACE BEFORE SUBMISSION')
box(s,.65,2.55,14.7,3.78); picture(s,results,.85,2.66,14.3,3.55)
box(s,.65,6.5,7.3,1.18); picture(s,trace,.83,6.53,6.95,1.1)
text(s,8.25,6.57,6.8,.9,'Lines = seed means; whiskers = min–max.\nHit probability = hits / all shots; ε = 10⁻⁶ relative.\nTrack A CVRP gap uses exact classical routing.',16,MUTED)
caption(s,'MOCK tuned reference-hit probabilities: S 18–25%, M 0–6%, L 0%; 1,000 shots/instance. Trace is separate schematic.')
note(s,'[85 秒] 这一页全部是 mock performance，只展示最终结果应该怎样组织，不是实测数据。左图是 feasibility rate，分母是所有 shots。中图是 best feasible clustering cost 相对 constrained clustering reference 的 model gap。右图把选出的 clusters exact-classically routed 后，与 full CVRP reference 比较。两个 gap 测量不同问题，不能互相代替。线是三个 seeds 的 mean，细点与 whiskers 显示范围，不是 confidence intervals。下面是独立的 schematic best-so-far curve，必须单调不增。Reference-hit probability 也必须用所有 shots 作分母，并明确 epsilon。实际提交前必须全部替换；没有 feasible shot 时 gap 应 undefined。不能仅凭这些 mock curves 声称 tuned pulse 优于 baseline。')

s=base(7,'What did not work','Balanced loads can cost better clusters.','Actual saved example: M_seed2 • both assignments are capacity-feasible • Q = 20')
box(s,.65,2.55,10.6,4.95); picture(s,failure,.85,2.8,10.2,4.4)
box(s,11.5,2.55,3.85,4.95)
text(s,11.8,2.95,3.2,.75,f'+{loss:.1f}%',36,CORAL,True)
text(s,11.8,3.82,3.15,.9,'clustering cost\nfrom balance bias',19)
text(s,11.8,5.0,3.15,1.8,'Next: isolate\nformulation loss,\nembedding loss,\nand sampling loss.',18,MUTED)
caption(s,'Takeaway: atom savings are real; feasibility and approximation costs must be measured. AI: drafting/figures/code; team reviews claims.')
note(s,'[50 秒] 最后是已经存在于 repository 的实际 limitation，不是 mock data。M_seed2 的 constrained optimum clustering cost 是 329.21，saved QUBO solution 的 clustering cost 是 374.60，高约 13.8%。两者都满足 capacity：前者 loads 是 20,8,19，后者是 13,15,19。QUBO 更喜欢后者，因为它的 loads 更均衡。这个例子说明：balance 不是 capacity，feasible uneven loads 也可能被多罚。另一个问题是它不保证任意 instance 都有 feasible ground state，所以一定要 validator 和 raw feasibility rate。我们的真实贡献目前是 atom reduction 与可解释的 formulation；接下来必须分别评估 formulation、embedding、sampling 的损失。AI 用于 draft、charts 和 code，所有公式和结论由团队 review。')

deck=OUT/'TrackA_7min_MOCK.pptx'; prs.save(deck)
metadata={'deck':str(deck),'mock_data':'mock_metrics.csv','real_inputs':['cvrp/qubo.py','instances/*.json','instances/weights.json'],
 'illustrative_layout':{'instance':inst.name,'coords_um':coords.tolist(),'pulser_register_valid':True,**fm},
 'illustrative_pulse':{'duration_ns':4000,'omega_peak_rad_per_us':10,'detuning_start':-30,'detuning_end':hmap.delta_global,'pulser_constructed':True,'emulated':False},
 'real_case':{'instance':case.name,'constrained_cost':refcost,'saved_qubo_solution_cost':qcost,'penalty_loss_pct':loss},
 'timings_seconds':[35,45,100,60,45,85,50]}
(OUT/'data_provenance.json').write_text(json.dumps(metadata,indent=2),encoding='utf-8')
print(json.dumps({'deck':str(deck),'slides':len(prs.slides),'timing_seconds':sum(metadata['timings_seconds']),'layout_eta':fm['eta'],'actual_penalty_loss_pct':loss},indent=2))
